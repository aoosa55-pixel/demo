import html as html_lib
import io
import json
import math
import os
import re
import signal
import threading
import time
import urllib.parse
import urllib.request
import datetime

import numpy as np
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from dotenv import load_dotenv
import openpyxl
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from st_aggrid import AgGrid, GridOptionsBuilder
from st_aggrid.shared import JsCode
from streamlit.runtime.scriptrunner import get_script_run_ctx

st.set_page_config(page_title="SO/PO Dashboard & Quotation (Demo)", layout="wide")

# --- ระบบสั่งปิด CMD อัตโนมัติเมื่อปิดหน้าเว็บ ---
def monitor_client_disconnect():
    ctx = get_script_run_ctx()
    if ctx is None:
        return
    
    session_id = ctx.session_id
    while True:
        time.sleep(2)
        from streamlit.runtime import Runtime
        if not Runtime.instance().is_active_session(session_id):
            os.kill(os.getpid(), signal.SIGTERM)
            break

# เปิดใช้เฉพาะตอนรันบนเครื่องตัวเอง (ตั้ง AUTO_SHUTDOWN=1 ใน .env) — ห้ามเปิดบน Streamlit Cloud
# ไม่งั้นเซิร์ฟเวอร์จะถูกปิดทันทีที่ผู้ใช้คนใดคนหนึ่งปิดหน้าเว็บ
load_dotenv()
if os.getenv("AUTO_SHUTDOWN") == "1" and 'disconnect_thread_started' not in st.session_state:
    st.session_state['disconnect_thread_started'] = True
    threading.Thread(target=monitor_client_disconnect, daemon=True).start()
# ------------------------------------------------

# ====================================================
# Custom CSS — Liquid Glass / iOS-inspired UI
# ====================================================
st.markdown("""
<style>
:root {
    --glass-bg: rgba(255, 255, 255, 0.58);
    --glass-bg-strong: rgba(255, 255, 255, 0.76);
    --glass-border: rgba(255, 255, 255, 0.78);
    --glass-shadow: 0 18px 55px rgba(75, 94, 150, 0.14);
    --ink: #20304d;
    --muted: #71809d;
    --blue: #527cf5;
    --violet: #9b7cf7;
}

.stApp {
    background:
        radial-gradient(circle at 8% 5%, rgba(145, 190, 255, .40), transparent 28%),
        radial-gradient(circle at 92% 8%, rgba(224, 181, 255, .36), transparent 30%),
        radial-gradient(circle at 65% 95%, rgba(156, 219, 255, .30), transparent 28%),
        linear-gradient(135deg, #eef5ff 0%, #f7f5ff 48%, #edf7ff 100%);
    color: var(--ink);
}

.stApp::before {
    content: "";
    position: fixed;
    inset: 0;
    pointer-events: none;
    background: linear-gradient(120deg, rgba(255,255,255,.18), transparent 42%);
    z-index: 0;
}

.main .block-container {
    position: relative;
    z-index: 1;
    max-width: 1600px;
    padding-top: 2.2rem;
    padding-bottom: 3rem;
}

h1, h2, h3, h4, p, label, .stMarkdown {
    color: var(--ink);
}
h1 {
    font-size: clamp(2rem, 3vw, 3.2rem) !important;
    font-weight: 800 !important;
    letter-spacing: -1.4px;
}
h2, h3 {
    font-weight: 750 !important;
    letter-spacing: -.5px;
}

div[data-baseweb="tab-highlight"],
div[data-baseweb="tab-border"] {
    display: none !important;
}

.stTabs [data-baseweb="tab-list"] {
    gap: 12px !important;
    background: rgba(255, 255, 255, 0.08) !important;
    border: 1.5px solid rgba(255, 255, 255, 0.6) !important;
    border-radius: 999px !important;
    padding: 8px 12px !important;
    backdrop-filter: blur(20px) saturate(200%) !important;
    -webkit-backdrop-filter: blur(20px) saturate(200%) !important;
    box-shadow: 
        0 8px 32px 0 rgba(0, 0, 0, 0.06),
        inset 0 1px 2px 0 rgba(255, 255, 255, 0.8) !important;
}

.stTabs [data-baseweb="tab"] {
    border-radius: 999px !important;
    padding: 10px 28px !important;
    color: #4a5568 !important;
    font-size: 15px !important;
    font-weight: 600 !important;
    letter-spacing: 0.3px !important;
    border: 1px solid transparent !important;
    background: transparent !important;
    transition: all 0.25s cubic-bezier(0.4, 0, 0.2, 1) !important;
}

.stTabs [data-baseweb="tab"]:hover {
    background: rgba(255, 255, 255, 0.25) !important;
    color: #1a202c !important;
}

.stTabs [aria-selected="true"] {
    color: #000000 !important;
    background: rgba(255, 255, 255, 0.45) !important;
    border: 1.5px solid rgba(255, 255, 255, 0.9) !important;
    border-radius: 999px !important;
    backdrop-filter: blur(15px) !important;
    box-shadow: 
        0 4px 16px rgba(31, 38, 135, 0.12),
        inset 0 1px 3px rgba(255, 255, 255, 1) !important;
}

div[data-testid="stExpander"],
div[data-testid="stVerticalBlockBorderWrapper"],
div[data-testid="stForm"] {
    background: var(--glass-bg);
    border: 1px solid var(--glass-border);
    border-radius: 26px;
    box-shadow: var(--glass-shadow);
    backdrop-filter: blur(24px) saturate(140%);
    -webkit-backdrop-filter: blur(24px) saturate(140%);
}
div[data-testid="stExpander"] {
    overflow: hidden;
    margin-bottom: 18px;
}
div[data-testid="stExpander"] details summary {
    padding: 14px 18px;
    font-weight: 750;
    color: var(--ink);
}

.stTextArea textarea,
.stTextInput input,
.stSelectbox [data-baseweb="select"],
.stNumberInput input {
    background: rgba(255,255,255,.68) !important;
    border: 1px solid rgba(255,255,255,.92) !important;
    border-radius: 17px !important;
    color: var(--ink) !important;
    box-shadow: inset 0 1px 0 rgba(255,255,255,.9), 0 5px 18px rgba(80,100,150,.06);
    backdrop-filter: blur(16px);
}
.stTextArea textarea:focus,
.stTextInput input:focus {
    border-color: #86a9ff !important;
    box-shadow: 0 0 0 3px rgba(105,145,255,.16) !important;
}

div.stButton > button,
.stDownloadButton button {
    border: 1px solid rgba(255,255,255,.9) !important;
    border-radius: 999px !important;
    padding: 10px 22px !important;
    font-weight: 750 !important;
    color: #34476b !important;
    background: rgba(255,255,255,.72) !important;
    box-shadow: 0 8px 22px rgba(83,104,160,.13), inset 0 1px 0 rgba(255,255,255,.95);
    transition: transform .18s ease, box-shadow .18s ease;
}
div.stButton > button:hover,
.stDownloadButton button:hover {
    transform: translateY(-2px);
    box-shadow: 0 12px 28px rgba(83,104,160,.22), inset 0 1px 0 rgba(255,255,255,1);
    color: #24385c !important;
}
div.stButton > button[kind="primary"] {
    color: white !important;
    border: none !important;
    background: linear-gradient(135deg, #4f8cff, #9b75f5) !important;
    box-shadow: 0 10px 25px rgba(103,126,245,.34) !important;
}

div[data-testid="stAlert"] {
    border-radius: 18px;
    border: 1px solid rgba(255,255,255,.75);
    box-shadow: 0 8px 24px rgba(80,100,150,.08);
    backdrop-filter: blur(18px);
}

hr {
    border-color: rgba(125, 145, 190, .20) !important;
}

div[data-testid="stAgGrid"] {
    padding: 10px;
    border-radius: 24px;
    background: rgba(255,255,255,.52);
    border: 1px solid rgba(255,255,255,.82);
    box-shadow: var(--glass-shadow);
    backdrop-filter: blur(22px);
}
.ag-theme-alpine {
    --ag-background-color: rgba(255,255,255,.20);
    --ag-foreground-color: #304363;
    --ag-header-background-color: rgba(111, 151, 245, .16);
    --ag-header-foreground-color: #304363;
    --ag-row-hover-color: rgba(126, 166, 255, .13);
    --ag-selected-row-background-color: rgba(126, 166, 255, .18);
    --ag-border-color: rgba(130, 150, 190, .16);
    --ag-row-border-color: rgba(130, 150, 190, .13);
    --ag-border-radius: 18px;
    --ag-wrapper-border-radius: 18px;
    font-family: "Segoe UI", system-ui, sans-serif;
}

#MainMenu { visibility: hidden; }
footer { visibility: hidden; }
header { background: transparent !important; }
</style>
""", unsafe_allow_html=True)


# ====================================================
# Helper Functions
# ====================================================
def parse_pct(val):
    if pd.isna(val) or val is None:
        return 0.0
    val_str = str(val).replace("%", "").strip()
    try:
        return float(val_str) / 100.0
    except ValueError:
        return 0.0


def calculate_custom_price(retail_price: float, category: str, cust_row: pd.Series, supplier_cost: float = 0.0):
    cat = str(category).upper()

    honda_big = parse_pct(cust_row.get("HONDA_BIG_BIKE"))
    honda_small = parse_pct(cust_row.get("HONDA_SMALL_BIKE"))
    yamaha_big = parse_pct(cust_row.get("YAMAHA_BIG_BIKE"))
    yamaha_small = parse_pct(cust_row.get("YAMAHA_SMALL_BIKE"))
    kawasaki_big = parse_pct(cust_row.get("KAWASAKI_BIG_BIKE"))
    kawasaki_small = parse_pct(cust_row.get("KAWASAKI_SMALL_BIKE"))

    net_markup = parse_pct(cust_row.get("NET_PRICE"))
    vat_rate = parse_pct(cust_row.get("VAT"))
    fob_rate = parse_pct(cust_row.get("FOB"))

    if "NET" in cat:
        base_cost = supplier_cost if supplier_cost > 0 else retail_price
        base_price = base_cost * (1.0 + net_markup)
        discount_label = f"NET (+{int(net_markup * 100)}%)"
    elif "HONDA" in cat:
        is_bigbike = "BIGBIKE" in cat or "BIG BIKE" in cat
        disc = honda_big if is_bigbike else honda_small
        base_price = retail_price * (1.0 - disc)
        discount_label = f"HONDA {'BIG' if is_bigbike else 'SMALL'} ({int(disc * 100)}%)"
    elif "YAMAHA" in cat:
        is_bigbike = "BIGBIKE" in cat or "BIG BIKE" in cat
        disc = yamaha_big if is_bigbike else yamaha_small
        base_price = retail_price * (1.0 - disc)
        discount_label = f"YAMAHA {'BIG' if is_bigbike else 'SMALL'} ({int(disc * 100)}%)"
    elif "KAWASAKI" in cat:
        is_bigbike = "BIGBIKE" in cat or "BIG BIKE" in cat
        disc = kawasaki_big if is_bigbike else kawasaki_small
        base_price = retail_price * (1.0 - disc)
        discount_label = f"KAWASAKI {'BIG' if is_bigbike else 'SMALL'} ({int(disc * 100)}%)"
    else:
        base_price = retail_price
        discount_label = "NET (0%)"

    final_raw_price = base_price * (1.0 + vat_rate + fob_rate)
    final_unit_price = int(math.ceil(round(final_raw_price, 4)))

    return final_unit_price, discount_label


def generate_excel_pi(df_items, cust_row):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PROFORMA INVOICE"
    ws.views.sheetView[0].showGridLines = True

    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_title = Font(name="Calibri", size=14, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), 
                         top=Side(style="thin"), bottom=Side(style="thin"))

    logo_filename = "Header_logo.png"
    header_logo_path = os.path.join(os.path.dirname(__file__), logo_filename)
    if not os.path.exists(header_logo_path):
        header_logo_path = logo_filename

    if os.path.exists(header_logo_path):
        img = Image(header_logo_path)
        img.width = 650
        img.height = 65
        ws.add_image(img, 'A1')

    ws.merge_cells("A6:G6")
    ws["A6"] = "PROFORMA INVOICE"
    ws["A6"].font = font_title
    ws["A6"].alignment = align_center

    today = datetime.datetime.now()
    today_str = today.strftime("%B %d, %Y").upper()
    pi_no = f"PI-{today.strftime('%y%m%d')}"

    customer_name = str(cust_row.get('CUSTOMER_NAME', '')).upper()
    contact_person = str(cust_row.get('CONTACT_PERSON', cust_row.get('ATTN', cust_row.get('ATTENTION', '')))).upper()
    if contact_person == 'NAN': contact_person = ''

    tel = str(cust_row.get('TEL', '')).upper()
    if tel == 'NAN': tel = ''

    destination = str(cust_row.get('DESTINATION', 'JAPAN')).upper()
    if destination == 'NAN': destination = 'JAPAN'
    
    address = str(cust_row.get('ADDRESS', '')).upper()
    if address == 'NAN': address = ''

    addr_line1 = address
    addr_line2 = ""
    if len(address) > 40:
        split_pos = address.find(',', len(address)//2 - 10)
        if split_pos == -1: split_pos = address.find(' ', len(address)//2)
        if split_pos != -1:
            addr_line1 = address[:split_pos+1].strip()
            addr_line2 = address[split_pos+1:].strip()

    left_info = [f"CONSIGNEE : {customer_name}", f"ADDRESS : {addr_line1}"]
    if addr_line2: left_info.append(addr_line2)
    if tel: left_info.append(f"TEL : {tel}")
    if contact_person: left_info.append(f"CONTACT PERSON : {contact_person}")

    def get_val(key, default_val):
        val = str(cust_row.get(key, default_val)).strip()
        return val if val.upper() != 'NAN' and val != '' else None

    right_info = [f"ISSUE DATE : {today_str}", f"NO. : {pi_no}"]

    term_val = get_val('TERM', 'EX-WORKS, BANGKOK')
    if term_val: right_info.append(f"TERM : {term_val.upper()}")

    payment_val = get_val('TERM_OF_PAYMENT', '50% WHEN CONFIRM ORDER, THE REST 50% BEFORE SHIPPING OUT')
    if payment_val:
        payment_str = payment_val.upper()
        if len(payment_str) > 35:
            split_pos = payment_str.find(',', len(payment_str)//2 - 10)
            if split_pos == -1: split_pos = payment_str.find(' ', len(payment_str)//2 - 5)
            if split_pos != -1:
                right_info.append(f"TERM OF PAYMENT : {payment_str[:split_pos+1].strip()}")
                right_info.append(payment_str[split_pos+1:].strip())
            else: right_info.append(f"TERM OF PAYMENT : {payment_str}")
        else: right_info.append(f"TERM OF PAYMENT : {payment_str}")

    from_val = get_val('FROM', 'BANGKOK , THAILAND')
    if from_val: right_info.append(f"FROM : {from_val.upper()}")

    to_val = get_val('TO', destination)
    if to_val: right_info.append(f"TO : {to_val.upper()}")

    origin_val = get_val('ORIGIN', 'THAILAND')
    if origin_val: right_info.append(f"ORIGIN : {origin_val.upper()}")

    max_rows = max(len(left_info), len(right_info))
    for i in range(max_rows):
        current_r = 8 + i
        if i < len(left_info):
            ws[f"A{current_r}"] = left_info[i]
            ws[f"A{current_r}"].font = font_regular
        if i < len(right_info):
            ws[f"E{current_r}"] = right_info[i]
            ws[f"E{current_r}"].font = font_regular

    start_row = 8 + max_rows + 1
    headers = ["NO.", "DESCRIPTION", "PARTS NUMBER", "QTY", "UNIT PRICE", "AMOUNT", "REMARKS"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = font_bold
        cell.alignment = align_center
        cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        cell.border = thin_border

    current_row = start_row + 1
    for idx, (_, row) in enumerate(df_items.iterrows(), 1):
        ws.cell(row=current_row, column=1, value=idx).alignment = align_center
        ws.cell(row=current_row, column=2, value=str(row.get("DESCRIPTION", "")).upper()).alignment = align_left
        ws.cell(row=current_row, column=3, value=str(row.get("PARTS NUMBER", "")).upper()).alignment = align_left

        qty_val = float(row.get("QTY", 0)) if pd.notna(row.get("QTY")) else 0
        qty_cell = ws.cell(row=current_row, column=4, value=int(qty_val))
        qty_cell.alignment = align_center
        qty_cell.number_format = '#,##0'

        price_val = float(row.get("UNIT PRICE", 0)) if pd.notna(row.get("UNIT PRICE")) else 0
        price_cell = ws.cell(row=current_row, column=5, value=price_val)
        price_cell.alignment = align_right
        price_cell.number_format = '#,##0.00'

        amount_cell = ws.cell(row=current_row, column=6, value=f"=D{current_row}*E{current_row}")
        amount_cell.alignment = align_right
        amount_cell.number_format = '#,##0.00'

        ws.cell(row=current_row, column=7, value=str(row.get("REMARKS", "")).upper()).alignment = align_left

        for c in range(1, 8):
            ws.cell(row=current_row, column=c).border = thin_border
            ws.cell(row=current_row, column=c).font = font_regular
        current_row += 1

    for c in range(1, 8):
        ws.cell(row=current_row, column=c).border = thin_border
    current_row += 1

    total_label = ws.cell(row=current_row, column=2, value="TOTAL")
    total_label.font = font_bold
    total_label.alignment = align_center

    qty_sum_cell = ws.cell(row=current_row, column=4, value=f"=SUM(D{start_row+1}:D{current_row-2})")
    qty_sum_cell.font = font_bold
    qty_sum_cell.alignment = align_center
    qty_sum_cell.number_format = '#,##0'

    total_cell = ws.cell(row=current_row, column=6, value=f"=SUM(F{start_row+1}:F{current_row-2})")
    total_cell.font = font_bold
    total_cell.alignment = align_right
    total_cell.number_format = '#,##0.00'

    bank_row = current_row + 2
    bank_details = [
        "BANK INFORMATION",
        "A/C NAME : MY MEGAPARTS CO.,LTD.",
        "A/C NUMBER : 735-1-03701-0",
        "BANK NAME : KASIKORN",
        "BANK BRANCH : SUKHABIBARN 3 BUNGKUM",
        "SWIFT CODE : KASITHBK"
    ]
    
    for i, text in enumerate(bank_details):
        cell = ws.cell(row=bank_row + i, column=1, value=text)
        cell.font = font_bold if i == 0 else font_regular
        cell.alignment = align_left

    sign_row = bank_row + 8
    ws.merge_cells(start_row=sign_row, start_column=2, end_row=sign_row, end_column=4)
    cust_sign_cell = ws.cell(row=sign_row, column=2, value="(..................................................)")
    cust_sign_cell.font = font_bold
    cust_sign_cell.alignment = align_center

    ws.merge_cells(start_row=sign_row+1, start_column=2, end_row=sign_row+1, end_column=4)
    cust_role_cell = ws.cell(row=sign_row+1, column=2, value="Customer's authorized signature and stamp")
    cust_role_cell.font = font_bold
    cust_role_cell.alignment = align_center

    ws.merge_cells(start_row=sign_row, start_column=5, end_row=sign_row, end_column=7)
    sign_cell = ws.cell(row=sign_row, column=5, value="( VISARUT PAKUMTIP )")
    sign_cell.font = font_bold
    sign_cell.alignment = align_center

    ws.merge_cells(start_row=sign_row+1, start_column=5, end_row=sign_row+1, end_column=7)
    role_cell = ws.cell(row=sign_row+1, column=5, value="ACCOUNT EXECUTIVE")
    role_cell.font = font_bold
    role_cell.alignment = align_center

    col_widths = {"A": 6, "B": 35, "C": 22, "D": 10, "E": 14, "F": 16, "G": 18}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_excel_full_quotation(df_items, cust_row):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QUOTATION"
    ws.views.sheetView[0].showGridLines = True

    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_title = Font(name="Calibri", size=14, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), 
                         top=Side(style="thin"), bottom=Side(style="thin"))

    logo_filename = "Header_logo.png"
    header_logo_path = os.path.join(os.path.dirname(__file__), logo_filename)
    if not os.path.exists(header_logo_path):
        header_logo_path = logo_filename

    if os.path.exists(header_logo_path):
        img = Image(header_logo_path)
        img.width = 650
        img.height = 65
        ws.add_image(img, 'A1')

    ws.merge_cells("A6:G6")
    ws["A6"] = "QUOTATION"
    ws["A6"].font = font_title
    ws["A6"].alignment = align_center

    today = datetime.datetime.now()
    today_str = today.strftime("%B %d, %Y").upper()
    q_no = f"Q{today.strftime('%y%m%d')}"

    customer_name = str(cust_row.get('CUSTOMER_NAME', '')).upper()
    contact_person = str(cust_row.get('CONTACT_PERSON', cust_row.get('ATTN', cust_row.get('ATTENTION', '')))).upper()
    if contact_person == 'NAN': contact_person = ''

    tel = str(cust_row.get('TEL', '')).upper()
    if tel == 'NAN': tel = ''

    destination = str(cust_row.get('DESTINATION', 'JAPAN')).upper()
    if destination == 'NAN': destination = 'JAPAN'
    
    address = str(cust_row.get('ADDRESS', '')).upper()
    if address == 'NAN': address = ''

    addr_line1 = address
    addr_line2 = ""
    if len(address) > 40:
        split_pos = address.find(',', len(address)//2 - 10)
        if split_pos == -1: split_pos = address.find(' ', len(address)//2)
        if split_pos != -1:
            addr_line1 = address[:split_pos+1].strip()
            addr_line2 = address[split_pos+1:].strip()

    left_info = [f"CONSIGNEE : {customer_name}", f"ADDRESS : {addr_line1}"]
    if addr_line2: left_info.append(addr_line2)
    if tel: left_info.append(f"TEL : {tel}")
    if contact_person: left_info.append(f"CONTACT PERSON : {contact_person}")

    def get_val(key, default_val):
        val = str(cust_row.get(key, default_val)).strip()
        return val if val.upper() != 'NAN' and val != '' else None

    right_info = [f"ISSUE DATE : {today_str}", f"NO. : {q_no}"]

    term_val = get_val('TERM', 'EX-WORKS, BANGKOK')
    if term_val: right_info.append(f"TERM : {term_val.upper()}")

    payment_val = get_val('TERM_OF_PAYMENT', '50% WHEN CONFIRM ORDER, THE REST 50% BEFORE SHIPPING OUT')
    if payment_val:
        payment_str = payment_val.upper()
        if len(payment_str) > 35:
            split_pos = payment_str.find(',', len(payment_str)//2 - 10)
            if split_pos == -1: split_pos = payment_str.find(' ', len(payment_str)//2 - 5)
            if split_pos != -1:
                right_info.append(f"TERM OF PAYMENT : {payment_str[:split_pos+1].strip()}")
                right_info.append(payment_str[split_pos+1:].strip())
            else: right_info.append(f"TERM OF PAYMENT : {payment_str}")
        else: right_info.append(f"TERM OF PAYMENT : {payment_str}")

    from_val = get_val('FROM', 'BANGKOK , THAILAND')
    if from_val: right_info.append(f"FROM : {from_val.upper()}")

    to_val = get_val('TO', destination)
    if to_val: right_info.append(f"TO : {to_val.upper()}")

    origin_val = get_val('ORIGIN', 'THAILAND')
    if origin_val: right_info.append(f"ORIGIN : {origin_val.upper()}")

    max_rows = max(len(left_info), len(right_info))
    for i in range(max_rows):
        current_r = 8 + i
        if i < len(left_info):
            ws[f"A{current_r}"] = left_info[i]
            ws[f"A{current_r}"].font = font_regular
        if i < len(right_info):
            ws[f"E{current_r}"] = right_info[i]
            ws[f"E{current_r}"].font = font_regular

    start_row = 8 + max_rows + 1
    headers = ["NO.", "DESCRIPTION", "PARTS NUMBER", "QTY", "UNIT PRICE", "AMOUNT", "REMARKS"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = font_bold
        cell.alignment = align_center
        cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        cell.border = thin_border

    current_row = start_row + 1
    for idx, (_, row) in enumerate(df_items.iterrows(), 1):
        ws.cell(row=current_row, column=1, value=idx).alignment = align_center
        ws.cell(row=current_row, column=2, value=str(row.get("DESCRIPTION", "")).upper()).alignment = align_left
        ws.cell(row=current_row, column=3, value=str(row.get("PARTS NUMBER", "")).upper()).alignment = align_left

        qty_val = float(row.get("QTY", 0)) if pd.notna(row.get("QTY")) else 0
        qty_cell = ws.cell(row=current_row, column=4, value=int(qty_val))
        qty_cell.alignment = align_center
        qty_cell.number_format = '#,##0'

        price_val = float(row.get("UNIT PRICE", 0)) if pd.notna(row.get("UNIT PRICE")) else 0
        price_cell = ws.cell(row=current_row, column=5, value=price_val)
        price_cell.alignment = align_right
        price_cell.number_format = '#,##0.00'

        amount_cell = ws.cell(row=current_row, column=6, value=f"=D{current_row}*E{current_row}")
        amount_cell.alignment = align_right
        amount_cell.number_format = '#,##0.00'

        ws.cell(row=current_row, column=7, value=str(row.get("REMARKS", "")).upper()).alignment = align_left

        for c in range(1, 8):
            ws.cell(row=current_row, column=c).border = thin_border
            ws.cell(row=current_row, column=c).font = font_regular
        current_row += 1

    for c in range(1, 8):
        ws.cell(row=current_row, column=c).border = thin_border
    current_row += 1

    total_label = ws.cell(row=current_row, column=2, value="TOTAL")
    total_label.font = font_bold
    total_label.alignment = align_center

    qty_sum_cell = ws.cell(row=current_row, column=4, value=f"=SUM(D{start_row+1}:D{current_row-2})")
    qty_sum_cell.font = font_bold
    qty_sum_cell.alignment = align_center
    qty_sum_cell.number_format = '#,##0'

    total_cell = ws.cell(row=current_row, column=6, value=f"=SUM(F{start_row+1}:F{current_row-2})")
    total_cell.font = font_bold
    total_cell.alignment = align_right
    total_cell.number_format = '#,##0.00'

    sign_row = current_row + 4

    ws.merge_cells(start_row=sign_row, start_column=5, end_row=sign_row, end_column=7)
    sign_cell = ws.cell(row=sign_row, column=5, value="( VISARUT PAKUMTIP )")
    sign_cell.font = font_bold
    sign_cell.alignment = align_center

    ws.merge_cells(start_row=sign_row+1, start_column=5, end_row=sign_row+1, end_column=7)
    role_cell = ws.cell(row=sign_row+1, column=5, value="ACCOUNT EXECUTIVE")
    role_cell.font = font_bold
    role_cell.alignment = align_center

    col_widths = {"A": 6, "B": 35, "C": 22, "D": 10, "E": 14, "F": 16, "G": 18}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def clean_part_no(value):
    if value is None:
        return ""
    return re.sub(r"[\s\u00a0\u200b]+", "", str(value).upper()).strip()


def add_part_names_to_csv(raw_input):
    csv_path = os.path.join(os.path.dirname(__file__), "eng_name.csv")
    parts = [clean_part_no(value) for value in re.split(r"[\n,]+", raw_input)]
    parts = list(dict.fromkeys(part for part in parts if part))

    if not os.path.exists(csv_path):
        pd.DataFrame(columns=["PARTS NUMBER", "DESCRIPTION", "MODELS"]).to_csv(csv_path, index=False, encoding="utf-8-sig")

    eng_names = pd.read_csv(csv_path, dtype=str, encoding="utf-8-sig").fillna("")
    existing = {clean_part_no(value) for value in eng_names["PARTS NUMBER"] if clean_part_no(value)}
    new_parts = [part for part in parts if part not in existing]
    duplicate_count = len(parts) - len(new_parts)
    if not new_parts:
        return 0, duplicate_count

    rows = []
    for part in new_parts:
        rows.append([part, f"PARTS {part}", "GENERAL"])

    with open(csv_path, "a", encoding="utf-8", newline="") as csv_file:
        for row in rows:
            csv_file.write(",".join('"' + value.replace('"', '""') + '"' for value in row) + "\n")
    return len(rows), duplicate_count


def build_invoice_data(df, customer_name, so_input):
    target_sos = {value.strip().upper() for value in re.split(r"[\n,]+", so_input) if value.strip()}
    selected = df[(df["customer_name"] == customer_name) & df["so_no"].fillna("").astype(str).str.strip().str.upper().isin(target_sos)].copy()
    if selected.empty:
        return pd.DataFrame()

    selected["remaining_packed"] = (selected["packed"] - selected["invoice_amount"]).clip(lower=0)
    selected["invoice_no"] = selected["invoice_no"].fillna("").astype(str).str.strip()
    selected = selected[~(selected["invoice_no"].ne("") & selected["remaining_packed"].eq(0) & selected["pending"].eq(0))]
    if selected.empty:
        return pd.DataFrame()

    rows = []
    for so_number, so_rows in selected.groupby("so_no", sort=False):
        rows.append({
            "is_header": True, "DESCRIPTION": so_number, "PARTSNUMBER": "", "QTY": "",
            "UNIT": "", "UNIT PRICE": "", "AMOUNT": "", "REMARKS": "",
            "PACKED": "", "PENDING": "", "INVOICE": "",
        })
        for _, row in so_rows.iterrows():
            quantity = int(row["qty"])
            rows.append({
                "is_header": False, "DESCRIPTION": row["description"], "PARTSNUMBER": row["part_no"],
                "QTY": quantity, "UNIT": row["unit"], "UNIT PRICE": row["unit_price"],
                "AMOUNT": quantity * row["unit_price"], "REMARKS": "",
                "PACKED": int(row["remaining_packed"]), "PENDING": int(row["pending"]), "INVOICE": row["invoice_no"],
            })
    return pd.DataFrame(rows)


def render_selection_calculator(dataframe):
    numeric_columns = ["QTY", "UNIT PRICE", "AMOUNT", "PACKED", "PENDING", "INVOICE AMOUNT"]
    centered_columns = {"QTY", "UNIT", "PACKED", "PENDING", "INVOICE AMOUNT", "INVOICE"}
    headers = [column for column in dataframe.columns if column not in ["is_header", "NO"]]
    component_height = min(760, max(300, 72 + (len(dataframe) * 34)))
    rows_html = []
    for row_index, (_, row) in enumerate(dataframe.iterrows()):
        if row["is_header"] and "SO" in headers:
            so_value = html_lib.escape(str(row["SO"]))
            rows_html.append(f'<tr><td class="so-header" colspan="{len(headers)}">{so_value}</td></tr>')
            continue
        cells = []
        for column_index, column in enumerate(headers):
            value = row[column]
            if column == "INVOICE" and str(value).strip().lower() in {"nan", "none"}:
                value = ""
            is_numeric = column in numeric_columns and not row["is_header"]
            cell_class = "numeric" if is_numeric else ""
            if column in centered_columns: cell_class += " centered"
            if column == "PACKED" and not row["is_header"] and str(value).replace(",", "") == "0":
                cell_class += " packed-zero"
            if row["is_header"]: cell_class += " so-header"
            numeric_value = str(value).replace(",", "") if is_numeric else ""
            cells.append(f'<td class="{cell_class}" data-row="{row_index}" data-col="{column_index}" data-value="{numeric_value}">{html_lib.escape(str(value))}</td>')
        rows_html.append("<tr>" + "".join(cells) + "</tr>")

    html = f"""
    <style>
        body {{ font-family: Arial, sans-serif; margin: 8px 8px 18px; }}
        table {{ border-collapse: collapse; width: 100%; table-layout: fixed; font-size: 12px; user-select: none; }}
        col {{ width: auto; }}
        th, td {{ border: 1px solid #d9dee5; padding: 7px 5px; min-height: 30px; white-space: normal; overflow-wrap: anywhere; line-height: 1.35; }}
        th {{ background: #f3f5f7; position: sticky; top: 0; z-index: 1; position: relative; }}
        th .resize-handle {{ position: absolute; top: 0; right: -3px; width: 7px; height: 100%; cursor: col-resize; z-index: 2; }}
        td {{ cursor: cell; }}
        td.numeric {{ text-align: right; }}
        th.centered, td.centered {{ text-align: center; }}
        td.so-header {{ background: #e9ecef; font-weight: bold; text-align: center; white-space: normal; overflow-wrap: anywhere; word-break: break-word; }}
        td.packed-zero {{ background: #fce8e6 !important; }}
        td.selected {{ background: #cfe8ff !important; outline: 1px solid #1976d2; }}
        #popup {{ display: none; position: fixed; z-index: 5; left: 50%; bottom: 10px; transform: translateX(-50%); background: white; border: 1px solid #8a99a8; box-shadow: 0 3px 12px #0003; padding: 8px; border-radius: 5px; }}
        #popup button {{ margin: 2px; padding: 5px 9px; cursor: pointer; }}
        #result {{ font-weight: bold; margin-left: 8px; }}
        .wrap {{ width: 100%; max-height: {component_height - 70}px; overflow: auto; padding-bottom: 8px; }}
    </style>
    <div class="wrap"><table id="data">
        <colgroup>{''.join('<col>' for _ in headers)}</colgroup>
        <thead><tr>{''.join(f'<th class="{"centered" if header in centered_columns else ""}">{html_lib.escape(header)}<span class="resize-handle"></span></th>' for header in headers)}</tr></thead>
        <tbody>{''.join(rows_html)}</tbody>
    </table></div>
    <div id="popup"><button onclick="calculate('sum')">SUM</button><button onclick="calculate('count')">COUNT</button><span id="result"></span></div>
    <script>
        let startCell = null, selectedCells = [], dragging = false, resizing = false;
        const table = document.getElementById('data'), popup = document.getElementById('popup');
        const columns = [...table.querySelectorAll('col')];
        document.querySelectorAll('.resize-handle').forEach((handle, index) => {{
            handle.addEventListener('mousedown', event => {{
                event.preventDefault(); event.stopPropagation(); resizing = true;
                const startX = event.clientX, startWidth = columns[index].getBoundingClientRect().width;
                function resize(moveEvent) {{ columns[index].style.width = Math.max(45, startWidth + moveEvent.clientX - startX) + 'px'; }}
                function finishResize() {{ resizing = false; document.removeEventListener('mousemove', resize); document.removeEventListener('mouseup', finishResize); }}
                document.addEventListener('mousemove', resize); document.addEventListener('mouseup', finishResize);
            }});
        }});
        function clearSelection() {{ document.querySelectorAll('.selected').forEach(cell => cell.classList.remove('selected')); }}
        function selectRectangle(endCell) {{
            clearSelection();
            const minRow = Math.min(+startCell.dataset.row, +endCell.dataset.row);
            const maxRow = Math.max(+startCell.dataset.row, +endCell.dataset.row);
            const minCol = Math.min(+startCell.dataset.col, +endCell.dataset.col);
            const maxCol = Math.max(+startCell.dataset.col, +endCell.dataset.col);
            selectedCells = [...table.querySelectorAll('td')].filter(cell =>
                +cell.dataset.row >= minRow && +cell.dataset.row <= maxRow &&
                +cell.dataset.col >= minCol && +cell.dataset.col <= maxCol);
            selectedCells.forEach(cell => cell.classList.add('selected'));
        }}
        table.addEventListener('mousedown', event => {{ const cell = event.target.closest('td'); if (!cell || resizing) return; event.preventDefault(); dragging = true; startCell = cell; popup.style.display = 'none'; selectRectangle(cell); }});
        table.addEventListener('mouseover', event => {{ if (dragging && event.target.closest('td')) selectRectangle(event.target.closest('td')); }});
        document.addEventListener('mouseup', event => {{ if (!dragging) return; dragging = false; if (selectedCells.length) {{ popup.style.display = 'block'; }} }});
        function calculate(mode) {{
            const values = selectedCells.map(cell => Number(cell.dataset.value)).filter(value => !Number.isNaN(value));
            const result = mode === 'sum' ? values.reduce((total, value) => total + value, 0) : values.length;
            document.getElementById('result').textContent = (mode === 'sum' ? 'SUM: ' : 'COUNT: ') + result.toLocaleString();
        }}
    </script>
    """
    components.html(html, height=component_height)


# ====================================================
# DB / Data Functions (Offline / Static Data Mode)
# ====================================================
@st.cache_data
def fetch_data():
    report_path = os.path.join(os.path.dirname(__file__), "odoo_po_so_report.xlsx")
    if not os.path.exists(report_path):
        report_path = os.path.join(os.path.dirname(__file__), "odoo_po_so_report.csv")

    if os.path.exists(report_path):
        df = pd.read_excel(report_path) if report_path.endswith(".xlsx") else pd.read_csv(report_path)
    else:
        st.error("❌ ไม่พบไฟล์ odoo_po_so_report.xlsx หรือ .csv สำหรับแสดงผลข้อมูล")
        return pd.DataFrame()

    # แปลงชื่อคอลัมน์ (ภาษาไทยจาก Odoo export) ให้เป็นชื่อมาตรฐานที่โค้ดใช้
    thai_col_map = {
        "เลข so": "so_no",
        "เลข po": "po_no",
        "ชื่อลูกค้า": "customer_name",
        "ชื่อคู่ค้า": "vendor_name",
        "รหัสสินค้า": "part_no",
        "ชื่อสินค้า": "description",
        "จำนวนสั่ง (po)": "qty",
        "จำนวนที่รับเข้า": "packed",
        "unit price": "unit_price",
        "เลข invoice (so)": "invoice_no",
        "วันที่ invoice (so)": "invoice_date",
        "จำนวนที่ตัด invoice แล้ว": "invoice_amount",
    }
    col_map = {}
    for c in df.columns:
        key = str(c).strip().lower()
        col_map[c] = thai_col_map.get(key, key)
    df.rename(columns=col_map, inplace=True)

    # ตรวจสอบและรองรับคอลัมน์สำคัญ
    required_cols = ["so_no", "po_no", "customer_name", "part_no", "description", "qty", "packed", "unit_price", "invoice_no", "invoice_amount"]
    for col in required_cols:
        if col not in df.columns:
            df[col] = 0 if col in ["qty", "packed", "unit_price", "invoice_amount"] else ""

    for col in ["so_no", "po_no", "customer_name", "part_no", "description", "invoice_no"]:
        df[col] = df[col].fillna("").astype(str).str.strip()
    df = df[df["so_no"] != ""].reset_index(drop=True)

    # แมปชื่อภาษาอังกฤษจาก eng_name.csv
    csv_path = os.path.join(os.path.dirname(__file__), "eng_name.csv")
    if os.path.exists(csv_path):
        eng_names = pd.read_csv(csv_path, dtype=str, encoding="utf-8-sig").fillna("")
        eng_names.columns = [str(c).strip().upper() for c in eng_names.columns]
        if {"PARTS NUMBER", "DESCRIPTION"}.issubset(eng_names.columns):
            eng_names["part_no_key"] = eng_names["PARTS NUMBER"].str.strip().str.upper()
            eng_names = eng_names.drop_duplicates("part_no_key")
            
            df["part_no_key"] = df["part_no"].fillna("").astype(str).str.strip().str.upper()
            df = df.merge(eng_names[["part_no_key", "DESCRIPTION"]], on="part_no_key", how="left")
            df["description"] = df["DESCRIPTION"].replace("", np.nan).fillna(df["description"])
            if "DESCRIPTION" in df.columns:
                df.drop(columns=["DESCRIPTION"], inplace=True)
            df.drop(columns=["part_no_key"], inplace=True)

    df["unit"] = "PCS"
    df["qty"] = pd.to_numeric(df["qty"], errors='coerce').fillna(0).round().astype(int)
    df["packed"] = pd.to_numeric(df["packed"], errors='coerce').fillna(0).round().astype(int)
    df["pending"] = df["qty"] - df["packed"]
    df["unit_price"] = pd.to_numeric(df["unit_price"], errors='coerce').fillna(0.0)
    df["amount"] = df["qty"] * df["unit_price"]
    df["invoice_amount"] = pd.to_numeric(df["invoice_amount"], errors='coerce').fillna(0).round().astype(int)

    return df


def fetch_live_quotation_prices(parsed_items):
    if not parsed_items:
        return pd.DataFrame()

    input_df = pd.DataFrame(parsed_items, columns=["PARTS NUMBER", "QTY"])
    input_df["PARTS NUMBER"] = input_df["PARTS NUMBER"].astype(str).str.strip().str.upper()

    csv_path = os.path.join(os.path.dirname(__file__), "eng_name.csv")
    if os.path.exists(csv_path):
        master_df = pd.read_csv(csv_path, dtype=str, encoding="utf-8-sig").fillna("")
        master_df.columns = [str(c).strip().upper() for c in master_df.columns]
        master_df["PARTS NUMBER"] = master_df["PARTS NUMBER"].str.strip().str.upper()
    else:
        master_df = pd.DataFrame(columns=["PARTS NUMBER", "DESCRIPTION", "MODELS", "CATEGORY", "RETAIL", "SUPPLIER_COST"])

    merged_df = input_df.merge(master_df, on="PARTS NUMBER", how="left")

    # eng_name.csv อาจไม่มีบางคอลัมน์ (เช่น RETAIL, SUPPLIER_COST) — สร้างไว้ก่อนเพื่อไม่ให้ error
    for col, default in [("DESCRIPTION", ""), ("MODELS", ""), ("CATEGORY", ""), ("RETAIL", 0.0), ("SUPPLIER_COST", 0.0)]:
        if col not in merged_df.columns:
            merged_df[col] = default

    merged_df["DESCRIPTION"] = merged_df["DESCRIPTION"].fillna("")
    category = merged_df["CATEGORY"].replace("", np.nan).fillna(merged_df["MODELS"].replace("", np.nan))
    merged_df["CATEGORY"] = category.fillna("NET")
    merged_df["VAT"] = "0%"
    merged_df["FOB"] = "0%"
    merged_df["RETAIL"] = pd.to_numeric(merged_df["RETAIL"], errors='coerce').fillna(0.0)
    merged_df["SUPPLIER_COST"] = pd.to_numeric(merged_df["SUPPLIER_COST"], errors='coerce').fillna(0.0)
    merged_df["UNIT PRICE"] = 0
    merged_df["DISCOUNT_BRAND"] = "-"
    merged_df["REMARKS"] = ""

    is_not_found = merged_df["DESCRIPTION"] == ""
    merged_df.loc[is_not_found, "REMARKS"] = "NOT FOUND IN DB"

    merged_df = merged_df[["DESCRIPTION", "PARTS NUMBER", "QTY", "CATEGORY", "VAT", "FOB", "RETAIL", "SUPPLIER_COST", "UNIT PRICE", "DISCOUNT_BRAND", "REMARKS"]]

    return merged_df


# ====================================================
# Streamlit Interface Layout
# ====================================================
st.title("SO/PO & Quotation Management System (Demo)")

tab1, tab2 = st.tabs(["📦 SO/PO Dashboard", "📄 Create Quotation"])

# ----------------------------------------------------
# TAB 1: SO/PO Dashboard 
# ----------------------------------------------------
with tab1:
    with st.expander("Add Part Names (English)"):
        part_input = st.text_area("Enter Part Numbers (comma or line-separated)", placeholder="e.g., 04801-K0G-900ZA, 06110-K0A-306", key="tab1_part_input")
        if st.button("Search & Save Parts", key="btn_save_parts"):
            if not part_input.strip():
                st.warning("Please enter at least 1 part number.")
            else:
                try:
                    added_count, duplicate_count = add_part_names_to_csv(part_input)
                    st.success(f"Added {added_count} new item(s) (Skipped {duplicate_count} duplicate(s)).")
                    fetch_data.clear()
                except Exception as error:
                    st.error(f"Failed to add parts: {error}")

    df_raw = fetch_data()
    
    if not df_raw.empty and "customer_name" in df_raw.columns:
        customer_list = [c for c in df_raw["customer_name"].dropna().unique().tolist() if str(c).strip() != ""]
        customer_list.sort()
    else:
        customer_list = []

    selected_customer = st.selectbox("🔍 Select Customer:", options=["-- Select Customer --"] + customer_list, key="tab1_customer_select")

    if selected_customer and selected_customer != "-- Select Customer --":
        df_filtered = df_raw[df_raw["customer_name"] == selected_customer].copy()

        if df_filtered.empty:
            st.warning("No data found for this customer.")
        else:
            if st.button("🔄 Refresh Data", key="btn_refresh_dashboard"):
                st.cache_data.clear()
                st.rerun()

            @st.dialog("Create Invoice by SO", width="large")
            def invoice_dialog(customer_name):
                invoice_so_input = st.text_area("Enter SO Numbers (comma or line-separated)", placeholder="e.g., SO001, SO002", key="invoice_so_input")
                if st.button("Generate Invoice Items", key="create_invoice_button"):
                    if not invoice_so_input.strip():
                        st.warning("Please enter at least 1 SO number.")
                    else:
                        invoice_df = build_invoice_data(df_raw, customer_name, invoice_so_input)
                        if invoice_df.empty:
                            st.warning("No items available for invoice or all items have been fully invoiced.")
                        else:
                            st.session_state["invoice_df"] = invoice_df

                if "invoice_df" in st.session_state:
                    invoice_df = st.session_state["invoice_df"]
                    render_selection_calculator(invoice_df)
                    output = io.BytesIO()
                    with pd.ExcelWriter(output, engine="openpyxl") as writer:
                        invoice_df.drop(columns=["is_header"]).to_excel(writer, index=False, sheet_name="Invoice")
                    st.download_button("Download Invoice Excel", data=output.getvalue(), file_name="invoice_items.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

            if st.button("Open Invoice Builder", key="open_invoice_dialog"):
                invoice_dialog(selected_customer)

            combined_rows = []
            so_list = df_filtered["so_no"].dropna().unique().tolist()
            item_counter = 1

            for so_key in so_list:
                df_so = df_filtered[df_filtered["so_no"] == so_key].reset_index(drop=True)
                combined_rows.append({"is_header": True, "NO": "", "DESCRIPTION": f"--- {so_key} ---", "PARTSNUMBER": "", "QTY": "", "UNIT": "", "UNIT PRICE": "", "AMOUNT": "", "REMARKS": "", "PACKED": "", "PENDING": "", "INVOICE AMOUNT": "", "INVOICE": ""})

                for idx, row in df_so.iterrows():
                    combined_rows.append({"is_header": False, "NO": str(item_counter), "DESCRIPTION": row["description"], "PARTSNUMBER": row["part_no"], "QTY": f"{row['qty']:,}", "UNIT": row["unit"], "UNIT PRICE": f"{row['unit_price']:,.2f}", "AMOUNT": f"{row['amount']:,.2f}", "REMARKS": "", "PACKED": row["packed"], "PENDING": f"{row['pending']:,}", "INVOICE AMOUNT": f"{row['invoice_amount']:,}", "INVOICE": row["invoice_no"] if row["invoice_no"] else ""})
                    item_counter += 1

            df_final = pd.DataFrame(combined_rows)
            render_selection_calculator(df_final)

# ----------------------------------------------------
# TAB 2: Quotation Generator
# ----------------------------------------------------
with tab2:
    st.header("🏷️ Real-time Pricing & Quotation Calculator (Demo)")
    
    raw_quote_parts = st.text_area(
        "📋 Input Part Numbers & Quantities (Format: Part#Qty or Part Qty):",
        placeholder="41201-MKY-E00 10\n41201-MKN-D50 #5\n45510-K73-T31",
        height=150,
        key="raw_quote_parts",
    )

    if st.button("🔍 Fetch Prices (Demo)", type="primary"):
        lines = [line.strip() for line in raw_quote_parts.splitlines() if line.strip()]
        parsed_items = []

        for line in lines:
            sub_parts = [p.strip() for p in line.strip().split(",") if p.strip()]
            for part_str in sub_parts:
                match = re.search(r"^([^\s#*xX,]+)[\s#*xX,]+(\d+)$", part_str)
                if match:
                    p_no = match.group(1).upper()
                    qty = int(match.group(2))
                else:
                    p_no = part_str.upper()
                    qty = 1
                parsed_items.append((p_no, qty))

        if parsed_items:
            st.session_state["quote_df"] = fetch_live_quotation_prices(parsed_items)
            st.success(f"Fetched {len(st.session_state['quote_df'])} item(s) successfully. Please select a customer discount profile below.")
        else:
            st.warning("Please enter at least 1 part number.")

    if "quote_df" in st.session_state and not st.session_state["quote_df"].empty:
        
        st.divider()
        
        csv_path = os.path.join(os.path.dirname(__file__), "customer_discounts.csv")
        if os.path.exists(csv_path):
            cust_df = pd.read_csv(csv_path, encoding="utf-8-sig").fillna("")
            cust_df.columns = [str(c).strip().upper() for c in cust_df.columns]
            target_col = "CUSTOMER_NAME" if "CUSTOMER_NAME" in cust_df.columns else ("CUSTOMER" if "CUSTOMER" in cust_df.columns else None)
            
            if target_col:
                cust_list = ["-- Select Customer Profile --"] + [c for c in cust_df[target_col].tolist() if str(c).strip() != ""]
            else:
                cust_list = ["-- Error: Column CUSTOMER_NAME Not Found in CSV --"]
        else:
            cust_df = pd.DataFrame()
            cust_list = ["-- customer_discounts.csv Not Found --"]

        selected_customer = st.selectbox(
            "👤 Select Customer Discount Profile:",
            options=cust_list,
            key="select_quotation_customer"
        )

        cust_row_dict = {}
        if selected_customer and selected_customer not in ["-- Select Customer Profile --", "-- customer_discounts.csv Not Found --", "-- Error: Column CUSTOMER_NAME Not Found in CSV --"]:
            target_col = "CUSTOMER_NAME" if "CUSTOMER_NAME" in cust_df.columns else "CUSTOMER"
            cust_row = cust_df[cust_df[target_col] == selected_customer].iloc[0]
            cust_row_dict = cust_row.to_dict()
            
            updated_prices = []
            updated_labels = []
            vat_vals = []
            fob_vals = []

            for _, row in st.session_state["quote_df"].iterrows():
                unit_price, label = calculate_custom_price(
                    retail_price=float(row.get("RETAIL", 0)),
                    category=str(row.get("CATEGORY", "")),
                    cust_row=cust_row,
                    supplier_cost=float(row.get("SUPPLIER_COST", 0))
                )
                updated_prices.append(unit_price)
                updated_labels.append(label)

                v_str = str(cust_row.get("VAT", "0%")).strip()
                f_str = str(cust_row.get("FOB", "0%")).strip()
                vat_vals.append(v_str if v_str and v_str != "nan" else "0%")
                fob_vals.append(f_str if f_str and f_str != "nan" else "0%")

            st.session_state["quote_df"]["UNIT PRICE"] = updated_prices
            st.session_state["quote_df"]["DISCOUNT_BRAND"] = updated_labels
            st.session_state["quote_df"]["VAT"] = vat_vals
            st.session_state["quote_df"]["FOB"] = fob_vals
            st.success(f"✅ Prices updated for **{selected_customer}** (Calculated with Net/Discounts/VAT/FOB and rounded up).")

        # ----------------------------------------------------
        # 1. Internal Calculator (AgGrid)
        # ----------------------------------------------------
        st.subheader("📝 Internal Calculator")

        df_source = st.session_state["quote_df"].copy()

        if "::auto_unique_id::" in df_source.columns:
            df_source = df_source.drop(columns=["::auto_unique_id::"])

        all_cols = list(df_source.columns)
        preferred_first = ["DESCRIPTION", "PARTS NUMBER", "QTY", "CATEGORY", "VAT", "FOB"]
        preferred_last = ["RETAIL", "UNIT PRICE", "DISCOUNT_BRAND", "REMARKS"]

        first_cols = [c for c in preferred_first if c in all_cols]
        last_cols = [c for c in preferred_last if c in all_cols]
        middle_cols = [c for c in all_cols if c not in first_cols and c not in last_cols]

        display_cols = first_cols + middle_cols + last_cols
        display_df = df_source[display_cols].copy()

        bigbike_style_js = JsCode("""
        function(params) {
            if (!params.value) return null;
            var val = params.value.toString().toUpperCase();
            if (val.includes("BIG BIKE") || val.includes("BIGBIKE") || (val.includes("BIG") && val.includes("BIKE"))) {
                return { 'backgroundColor': '#FFE6E6', 'color': '#CC0000', 'fontWeight': 'bold' };
            }
            return null;
        }
        """)

        gb = GridOptionsBuilder.from_dataframe(display_df)
        gb.configure_default_column(resizable=True)
        gb.configure_column("::auto_unique_id::", hide=True)

        for col in display_cols:
            col_upper = str(col).upper()
            if col in ["QTY", "UNIT PRICE", "REMARKS"]:
                gb.configure_column(col, editable=True)
            else:
                gb.configure_column(col, editable=False)
                
            if any(k in col_upper for k in ["QTY", "PRICE", "RETAIL", "AMOUNT"]):
                gb.configure_column(col, type=["numericColumn"])

            if col in ["CATEGORY", "DISCOUNT_BRAND"]:
                gb.configure_column(col, cellStyle=bigbike_style_js)

        gridOptions = gb.build()

        ag_response = AgGrid(
            display_df,
            gridOptions=gridOptions,
            update_on=['rowDataUpdated', 'cellEditingStopped'],
            allow_unsafe_jscode=True,
            theme="alpine",
            height=300,
            enable_enterprise_modules=False,
        )

        edited_df = pd.DataFrame(ag_response["data"])

        if "::auto_unique_id::" in edited_df.columns:
            edited_df = edited_df.drop(columns=["::auto_unique_id::"])

        # ----------------------------------------------------
        # 2. Export Excel Files
        # ----------------------------------------------------
        date_str = datetime.datetime.now().strftime("%y%m%d")
        q_number = f"Q{date_str}"

        customer_name_clean = "Unknown_Customer"
        if selected_customer and selected_customer not in ["-- Select Customer Profile --", "-- customer_discounts.csv Not Found --", "-- Error: Column CUSTOMER_NAME Not Found in CSV --"]:
            customer_name_clean = str(selected_customer).replace('/', '_').replace('\\', '_').strip()

        file_name = f"QT_{customer_name_clean}_{q_number}.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = q_number

        headers = [
            "NO.", "DESCRIPTION", "PARTS NUMBER", "QTY", "UNIT PRICE", 
            "AMOUNT", "REMARKS", "", "RETAIL", "CATEGORY", "DISCOUNT_BRAND"
        ]
        ws.append(headers)

        header_font = Font(bold=True)
        header_align = Alignment(horizontal="center", vertical="center")
        center_align = Alignment(horizontal="center", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        red_font = Font(color="FF0000")
        light_red_fill = PatternFill(start_color="FFE6E6", end_color="FFE6E6", fill_type="solid")
        
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        for col_idx, cell in enumerate(ws[1], 1):
            cell.font = header_font
            cell.alignment = header_align
            if col_idx <= 7:
                cell.border = thin_border

        for idx, row in edited_df.iterrows():
            qty_val = float(row.get("QTY", 0)) if pd.notna(row.get("QTY")) else 0
            unit_price_val = float(row.get("UNIT PRICE", 0)) if pd.notna(row.get("UNIT PRICE")) else 0
            r_idx = idx + 2

            row_data = [
                idx + 1,
                row.get("DESCRIPTION", ""),
                row.get("PARTS NUMBER", ""),
                qty_val,
                unit_price_val,
                f"=D{r_idx}*E{r_idx}",
                row.get("REMARKS", ""),
                "",
                row.get("RETAIL", 0),
                row.get("CATEGORY", ""),
                row.get("DISCOUNT_BRAND", "")
            ]
            ws.append(row_data)
            current_row = ws.max_row
            
            for c in range(1, 8):
                ws.cell(row=current_row, column=c).border = thin_border
            
            ws.cell(row=current_row, column=1).alignment = center_align
            ws.cell(row=current_row, column=4).alignment = center_align
            ws.cell(row=current_row, column=7).alignment = center_align
            
            ws.cell(row=current_row, column=5).alignment = right_align
            ws.cell(row=current_row, column=6).alignment = right_align
            ws.cell(row=current_row, column=9).alignment = right_align

            ws.cell(row=current_row, column=7).font = red_font
            
            category_text = str(row.get("CATEGORY", "")).upper()
            if ("BIG" in category_text and "BIKE" in category_text) or ("BIGBIKE" in category_text.replace(" ", "")):
                ws.cell(row=current_row, column=10).fill = light_red_fill
                ws.cell(row=current_row, column=11).fill = light_red_fill

            ws.cell(row=current_row, column=4).number_format = '#,##0'
            ws.cell(row=current_row, column=5).number_format = '#,##0.00'
            ws.cell(row=current_row, column=6).number_format = '#,##0.00'
            ws.cell(row=current_row, column=9).number_format = '#,##0.00'

        last_data_row = len(edited_df) + 1

        blank_row = last_data_row + 1
        ws.append([""] * len(headers))
        for c in range(1, 8):
            ws.cell(row=blank_row, column=c).border = thin_border

        sum_row = blank_row + 1

        qty_sum_cell = ws.cell(row=sum_row, column=4, value=f"=SUM(D2:D{last_data_row})")
        amt_sum_cell = ws.cell(row=sum_row, column=6, value=f"=SUM(F2:F{last_data_row})")
        q_num_cell = ws.cell(row=sum_row, column=7, value=q_number)

        qty_sum_cell.font = Font(bold=True)
        qty_sum_cell.alignment = center_align
        qty_sum_cell.number_format = '#,##0'

        amt_sum_cell.font = Font(bold=True)
        amt_sum_cell.alignment = right_align
        amt_sum_cell.number_format = '#,##0.00'

        q_num_cell.font = Font(bold=True)
        q_num_cell.alignment = center_align

        column_widths = {
            'A': 6, 'B': 40, 'C': 20, 'D': 8, 'E': 12, 
            'F': 14, 'G': 18, 'H': 3, 'I': 12, 'J': 35, 'K': 25
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        excel_qt = excel_buffer.getvalue()

        excel_full_qt = generate_excel_full_quotation(edited_df, cust_row_dict)
        excel_pi = generate_excel_pi(edited_df, cust_row_dict)

        st.markdown("<br>", unsafe_allow_html=True)
        col1, col2, col3 = st.columns(3)
        
        with col1:
            st.download_button(
                label="📄 Simple Quotation",
                data=excel_qt,
                file_name=file_name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                type="primary"
            )

        with col2:
            st.download_button(
                label="📄 Full Quotation",
                data=excel_full_qt,
                file_name=f"QT_FULL_{customer_name_clean}_{date_str}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                type="primary"
            )

        with col3:
            st.download_button(
                label="📝 Proforma Invoice (PI)",
                data=excel_pi,
                file_name=f"PI_{customer_name_clean}_{date_str}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                type="primary"
            )


# ====================================================
# Final rounded-corner overrides
# ====================================================
st.markdown("""
<style>
.stTabs [data-baseweb="tab-list"] {
    border-radius: 999px !important;
    overflow: hidden !important;
}

.stTabs [data-baseweb="tab"] {
    border-radius: 999px !important;
    overflow: hidden !important;
}

.stTabs [data-baseweb="tab-highlight"] {
    border-radius: 999px !important;
}

div[data-baseweb="select"],
div[data-baseweb="select"] > div,
div[data-testid="stSelectbox"] > div,
div[data-testid="stSelectbox"] [role="combobox"] {
    border-radius: 999px !important;
}

div[data-testid="stSelectbox"] [role="combobox"] {
    background: rgba(255,255,255,.68) !important;
    border: 1px solid rgba(255,255,255,.92) !important;
    box-shadow: inset 0 1px 0 rgba(255,255,255,.9),
                0 5px 18px rgba(80,100,150,.06) !important;
    overflow: hidden !important;
}

div[data-testid="stTextArea"] textarea,
div[data-testid="stTextInput"] input,
div[data-testid="stNumberInput"] input {
    border-radius: 22px !important;
}

div[data-testid="stExpander"],
div[data-testid="stExpander"] details,
div[data-testid="stExpander"] summary {
    border-radius: 26px !important;
}

div[data-testid="stAlert"] {
    border-radius: 22px !important;
    overflow: hidden !important;
}

div.stButton > button,
.stDownloadButton button {
    border-radius: 999px !important;
}

div[data-testid="stDataFrame"],
div[data-testid="stAgGrid"],
div[data-testid="stTable"] {
    border-radius: 26px !important;
    overflow: hidden !important;
}

div[data-testid="stDataFrame"] > div,
div[data-testid="stAgGrid"] > div {
    border-radius: 24px !important;
    overflow: hidden !important;
}

.stTabs [data-baseweb="tab"]:focus-visible,
button:focus-visible,
input:focus-visible,
textarea:focus-visible {
    outline: 3px solid rgba(105,145,255,.35) !important;
    outline-offset: 2px !important;
}
</style>
""", unsafe_allow_html=True)